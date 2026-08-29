from pathlib import Path

import pytest
from docx import Document

from resume_matcher.document_loader import (
    DocumentExtractionError,
    extract_document_text,
)


class FakePdfPage:
    def __init__(
        self,
        text: str | None,
    ) -> None:
        self.text = text

    def extract_text(self) -> str | None:
        return self.text


class FakePdfReader:
    def __init__(
        self,
        path: Path,
    ) -> None:
        self.is_encrypted = False
        self.pages = [
            FakePdfPage("First PDF page"),
            FakePdfPage(None),
            FakePdfPage("Second PDF page"),
        ]


class EmptyPdfReader:
    def __init__(
        self,
        path: Path,
    ) -> None:
        self.is_encrypted = False
        self.pages = [
            FakePdfPage(None),
            FakePdfPage("   "),
        ]


def test_extracts_utf8_txt_text(
    tmp_path: Path,
) -> None:
    path = tmp_path / "resume.txt"

    # Include a UTF-8 byte-order mark to verify utf-8-sig handling.
    path.write_bytes(
        b"\xef\xbb\xbfPython developer"
    )

    assert extract_document_text(path) == "Python developer"


def test_extracts_text_from_pdf_pages(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "resume.pdf"
    path.write_bytes(b"fake PDF contents")

    monkeypatch.setattr(
        "resume_matcher.document_loader.PdfReader",
        FakePdfReader,
    )

    assert extract_document_text(path) == (
        "First PDF page\n\nSecond PDF page"
    )


def test_rejects_pdf_without_readable_text(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "scanned.pdf"
    path.write_bytes(b"fake scanned PDF")

    monkeypatch.setattr(
        "resume_matcher.document_loader.PdfReader",
        EmptyPdfReader,
    )

    with pytest.raises(
        DocumentExtractionError,
        match="may be scanned and require OCR",
    ):
        extract_document_text(path)


def test_extracts_docx_paragraphs_and_tables(
    tmp_path: Path,
) -> None:
    path = tmp_path / "resume.docx"

    document = Document()
    document.add_paragraph("Software engineer")

    table = document.add_table(
        rows=1,
        cols=2,
    )
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "Machine learning"

    document.save(path)

    assert extract_document_text(path) == (
        "Software engineer\n"
        "Python | Machine learning"
    )


def test_rejects_unsupported_document_type(
    tmp_path: Path,
) -> None:
    path = tmp_path / "resume.rtf"
    path.write_text(
        "Resume content",
        encoding="utf-8",
    )

    with pytest.raises(
        DocumentExtractionError,
        match="Unsupported document type",
    ):
        extract_document_text(path)


def test_rejects_missing_document(
    tmp_path: Path,
) -> None:
    path = tmp_path / "missing.pdf"

    with pytest.raises(
        FileNotFoundError,
        match="Document does not exist",
    ):
        extract_document_text(path)