from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
}


class DocumentExtractionError(ValueError):
    """Raised when readable text cannot be extracted from a document."""


def extract_document_text(
    file_path: str | Path,
) -> str:
    """Extract text from a TXT, PDF, or DOCX document."""

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Document does not exist: {path}"
        )

    if not path.is_file():
        raise DocumentExtractionError(
            f"Document path is not a file: {path}"
        )

    extension = path.suffix.casefold()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError(
            f"Unsupported document type: "
            f"{extension or 'no extension'}. "
            f"Supported types: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    if extension == ".txt":
        text = _extract_txt_text(path)
    elif extension == ".pdf":
        text = _extract_pdf_text(path)
    else:
        text = _extract_docx_text(path)

    text = text.strip()

    # A PDF containing only scanned images has no embedded text for
    # pypdf to extract. OCR support can be added in a future feature.
    if not text:
        if extension == ".pdf":
            raise DocumentExtractionError(
                "No readable text was found in the PDF. "
                "The PDF may be scanned and require OCR."
            )

        raise DocumentExtractionError(
            "No readable text was found in the document."
        )

    return text


def _extract_txt_text(
    path: Path,
) -> str:
    """Read a UTF-8 plain-text document."""

    try:
        # utf-8-sig also removes a UTF-8 byte-order mark when present.
        return path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentExtractionError(
            "The TXT document is not valid UTF-8 text."
        ) from error
    except OSError as error:
        raise DocumentExtractionError(
            f"Could not read TXT document: {path}"
        ) from error


def _extract_pdf_text(
    path: Path,
) -> str:
    """Extract embedded text from every PDF page."""

    try:
        reader = PdfReader(path)

        if reader.is_encrypted:
            raise DocumentExtractionError(
                "Encrypted PDF documents are not supported."
            )

        page_texts = [
            text.strip()
            for page in reader.pages
            if (text := page.extract_text())
            and text.strip()
        ]

        # Separate pages so their final and first lines do not merge.
        return "\n\n".join(page_texts)

    except DocumentExtractionError:
        raise
    except (PdfReadError, OSError) as error:
        raise DocumentExtractionError(
            f"Could not read PDF document: {path}"
        ) from error


def _extract_docx_text(
    path: Path,
) -> str:
    """Extract text from DOCX paragraphs and tables."""

    try:
        document = Document(path)

        text_parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Some résumés use tables for columns, skills, or contact details.
        for table in document.tables:
            for row in table.rows:
                cell_texts = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cell_texts:
                    text_parts.append(
                        " | ".join(cell_texts)
                    )

        return "\n".join(text_parts)

    except (
        PackageNotFoundError,
        BadZipFile,
        OSError,
    ) as error:
        raise DocumentExtractionError(
            f"Could not read DOCX document: {path}"
        ) from error